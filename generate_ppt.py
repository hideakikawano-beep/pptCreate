#!/usr/bin/env python3
"""
文字起こしデータからビジネスデザイン仕様のPowerPointを自動生成するスクリプト

デザイン仕様:
  - 配色: #2180FF (Blue) + #262626 (Soft Black)
  - 全スライド共通: Header (Title 16pt Blue / Separator / Key Message 20pt Bold) + Body (Cards)
  - カード: シャープコーナー (0px)、ボーダーまたはライトグレー塗り
  - フォント: Noto Sans
  - レイアウト自動分岐: 4件→2x2 Grid / 3件以下→Horizontal Row / 5件以上→Dynamic

使い方:
    python generate_ppt.py [transcription_file]

環境変数:
    ANTHROPIC_API_KEY: Claude API キー（必須）
"""

import sys
import json
import os
import re
from pathlib import Path

import anthropic
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# ─── カラー定数 ───────────────────────────────────────────────────────────────

BLUE            = RGBColor(0x21, 0x80, 0xFF)   # #2180FF
SOFT_BLACK      = RGBColor(0x26, 0x26, 0x26)   # #262626
WHITE           = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY      = RGBColor(0xE0, 0xE0, 0xE0)   # #E0E0E0  (border)
FILL_GRAY       = RGBColor(0xF5, 0xF5, 0xF5)   # #F5F5F5  (filled card)
FILL_BLUE_LIGHT = RGBColor(0xEB, 0xF5, 0xFF)   # #EBF5FF  (blue-tinted card)
SEPARATOR_CLR   = RGBColor(0xD0, 0xD0, 0xD0)   # separator line


# ─── スライド寸法 ─────────────────────────────────────────────────────────────

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
MARGIN  = Inches(0.6)

# ヘッダー領域
TITLE_TOP    = Inches(0.35)
TITLE_H      = Inches(0.45)
SEP_TOP      = TITLE_TOP + TITLE_H + Inches(0.12)
SEP_H        = Inches(0.015)
KM_TOP       = SEP_TOP + SEP_H + Inches(0.14)
KM_H         = Inches(0.75)

# ボディ領域
BODY_TOP     = KM_TOP + KM_H + Inches(0.25)
BODY_H       = SLIDE_H - BODY_TOP - Inches(0.35)
BODY_W       = SLIDE_W - MARGIN * 2
CARD_GAP     = Inches(0.18)
CARD_PAD     = Inches(0.18)


# ─── 共通ユーティリティ ───────────────────────────────────────────────────────

def _set_bg_white(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def _add_rect(slide, left, top, width, height,
              fill_color=None, border_color=None, border_pt=0.75):
    """シャープコーナー矩形を追加"""
    shape = slide.shapes.add_shape(1, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_pt)
    else:
        shape.line.fill.background()
    return shape


def _add_text_box(slide, text, left, top, width, height,
                  font_size, bold=False, color=SOFT_BLACK,
                  align=PP_ALIGN.LEFT, word_wrap=True):
    """テキストボックスを追加 (Noto Sans)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Noto Sans"
    return txBox


# ─── ヘッダー（全スライド共通） ──────────────────────────────────────────────

def _draw_header(slide, slide_title: str, key_message: str):
    """
    Title     → Blue (#2180FF), Normal, 16pt
    ─ separator ─
    Key Message → Soft Black (#262626), Bold, 20pt
    """
    content_w = SLIDE_W - MARGIN * 2

    # タイトル
    _add_text_box(slide, slide_title,
                  left=MARGIN, top=TITLE_TOP,
                  width=content_w, height=TITLE_H,
                  font_size=16, bold=False, color=BLUE)

    # セパレーター線
    _add_rect(slide,
              left=MARGIN, top=SEP_TOP,
              width=content_w, height=SEP_H,
              fill_color=SEPARATOR_CLR)

    # Key Message
    _add_text_box(slide, key_message,
                  left=MARGIN, top=KM_TOP,
                  width=content_w, height=KM_H,
                  font_size=20, bold=True, color=SOFT_BLACK)


# ─── カード ──────────────────────────────────────────────────────────────────

def _draw_card(slide, left, top, width, height,
               heading: str, body: str, style="filled"):
    """
    style="filled"   → Light Gray fill (#F5F5F5), no border
    style="bordered" → White fill + Light Gray border (#E0E0E0)
    """
    if style == "bordered":
        _add_rect(slide, left, top, width, height,
                  fill_color=WHITE, border_color=LIGHT_GRAY, border_pt=0.75)
    else:
        _add_rect(slide, left, top, width, height,
                  fill_color=FILL_GRAY)

    # 見出し: H1 — 18pt Bold Blue
    if heading:
        _add_text_box(slide, heading,
                      left=left + CARD_PAD, top=top + CARD_PAD,
                      width=width - CARD_PAD * 2, height=Inches(0.38),
                      font_size=18, bold=True, color=BLUE)

    # 本文: 12pt Normal
    if body:
        body_top_offset = (CARD_PAD + Inches(0.42)) if heading else CARD_PAD
        remaining_h = height - body_top_offset - CARD_PAD
        _add_text_box(slide, body,
                      left=left + CARD_PAD, top=top + body_top_offset,
                      width=width - CARD_PAD * 2, height=remaining_h,
                      font_size=12, bold=False, color=SOFT_BLACK)


# ─── ボディレイアウト（レイアウト自動分岐） ──────────────────────────────────

def _draw_body(slide, items: list):
    """
    items: [{"heading": str, "body": str}, ...]

    レイアウト分岐:
      4件         → 2x2 Grid  (bordered style)
      3件以下      → Horizontal Row (filled style)
      5件以上      → Dynamic 2-column Grid (filled style)
    """
    n = len(items)
    if n == 0:
        return

    if n <= 3:
        # ── Horizontal Row ──
        card_w = (BODY_W - CARD_GAP * (n - 1)) / n
        for i, item in enumerate(items):
            left = MARGIN + i * (card_w + CARD_GAP)
            _draw_card(slide, left, BODY_TOP, card_w, BODY_H,
                       item.get("heading", ""), item.get("body", ""),
                       style="filled")

    elif n == 4:
        # ── 2x2 Grid (MUST) ──
        card_w = (BODY_W - CARD_GAP) / 2
        card_h = (BODY_H - CARD_GAP) / 2
        positions = [
            (MARGIN,                    BODY_TOP),
            (MARGIN + card_w + CARD_GAP, BODY_TOP),
            (MARGIN,                    BODY_TOP + card_h + CARD_GAP),
            (MARGIN + card_w + CARD_GAP, BODY_TOP + card_h + CARD_GAP),
        ]
        for item, (l, t) in zip(items, positions):
            _draw_card(slide, l, t, card_w, card_h,
                       item.get("heading", ""), item.get("body", ""),
                       style="bordered")

    else:
        # ── Dynamic Layout (2-column grid) ──
        cols = 2
        rows = (n + 1) // 2
        card_w = (BODY_W - CARD_GAP) / cols
        card_h = (BODY_H - CARD_GAP * (rows - 1)) / rows
        for i, item in enumerate(items):
            col = i % cols
            row = i // cols
            left = MARGIN + col * (card_w + CARD_GAP)
            top  = BODY_TOP + row * (card_h + CARD_GAP)
            _draw_card(slide, left, top, card_w, card_h,
                       item.get("heading", ""), item.get("body", ""),
                       style="filled")


# ─── スライドビルダー ─────────────────────────────────────────────────────────

def _build_title_slide(prs: Presentation, data: dict):
    """タイトルスライド（左側 Blue アクセントバー）"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg_white(slide)

    # 左アクセントバー
    _add_rect(slide,
              left=Inches(0), top=Inches(0),
              width=Inches(0.14), height=SLIDE_H,
              fill_color=BLUE)

    # メインタイトル
    _add_text_box(slide, data.get("title", ""),
                  left=Inches(0.8), top=Inches(2.1),
                  width=SLIDE_W - Inches(1.4), height=Inches(1.5),
                  font_size=32, bold=True, color=SOFT_BLACK)

    # セパレーター
    _add_rect(slide,
              left=Inches(0.8), top=Inches(3.7),
              width=Inches(5.0), height=Inches(0.04),
              fill_color=BLUE)

    # サブタイトル
    subtitle = data.get("content", data.get("subtitle", ""))
    if subtitle:
        _add_text_box(slide, subtitle,
                      left=Inches(0.8), top=Inches(3.85),
                      width=SLIDE_W - Inches(1.4), height=Inches(0.7),
                      font_size=16, bold=False,
                      color=RGBColor(0x60, 0x60, 0x60))


def _build_content_slide(prs: Presentation, data: dict):
    """通常コンテンツスライド: Header + Body"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg_white(slide)

    # ヘッダー
    _draw_header(
        slide,
        slide_title=data.get("slide_title", data.get("title", "")),
        key_message=data.get("key_message", ""),
    )

    # ボディ
    body_items = data.get("body_items", [])
    if not body_items and data.get("bullets"):
        body_items = [{"heading": "", "body": b} for b in data["bullets"]]
    _draw_body(slide, body_items)

    # 発表者メモ
    notes = data.get("notes", "")
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


# ─── Claude API: リサーチ＆構成作成 ──────────────────────────────────────────

SYSTEM_PROMPT = """あなたは世界最高峰のビジネスデザイナーです。
文字起こしデータとウェブリサーチをもとに、以下の厳密なデザインルールに従った
プレゼンテーション構成を JSON で返してください。

【各スライドの構造】
- slide_title  : スライドタイトル（16pt Blue、簡潔に）
- key_message  : このスライドの結論・要点（20pt Bold、1〜2行）
- body_items   : カードリスト（各カードは heading + body）

【body_items 件数とレイアウトの対応】
- 4件   → 2x2 Grid（必須）
- 3件以下 → Horizontal Row
- 5件以上 → Dynamic 2列グリッド

【出力 JSON 形式（これ以外のテキスト不可）】
{
  "title": "プレゼンタイトル",
  "subtitle": "サブタイトル / 日付など",
  "slides": [
    {
      "type": "title",
      "title": "メインタイトル",
      "content": "サブタイトルや概要"
    },
    {
      "type": "content",
      "slide_title": "スライドタイトル",
      "key_message": "このスライドの結論（1〜2行）",
      "body_items": [
        {"heading": "カード見出し", "body": "カード本文（統計・具体例含む）"},
        {"heading": "カード見出し", "body": "カード本文"}
      ],
      "notes": "発表者メモ"
    }
  ]
}"""


def research_and_structure(transcription: str) -> dict:
    """Claude + Web検索でリサーチ＆プレゼン構成をJSONで返す"""
    client = anthropic.Anthropic()

    messages = [{
        "role": "user",
        "content": (
            "以下の文字起こしデータを分析し、ウェブ検索で関連情報・統計を補完して、"
            "プロフェッショナルなプレゼンテーション構成（8〜12スライド）をJSONで返してください。\n\n"
            f"## 文字起こしデータ:\n{transcription}"
        ),
    }]

    print("Claude がリサーチ＆構成作成中...", flush=True)

    with client.messages.stream(
        model="claude-opus-4-6",
        max_tokens=8000,
        thinking={"type": "adaptive"},
        system=SYSTEM_PROMPT,
        tools=[{"type": "web_search_20260209", "name": "web_search"}],
        messages=messages,
    ) as stream:
        current = None
        for event in stream:
            if not hasattr(event, "type"):
                continue
            if event.type == "content_block_start" and hasattr(event, "content_block"):
                current = event.content_block.type
                if current == "thinking":
                    print("  [思考中...]", flush=True)
                elif (current == "tool_use"
                      and getattr(event.content_block, "name", "") == "web_search"):
                    print("  [ウェブ検索中...]", flush=True)
                elif current == "text":
                    print("  [JSON生成中", end="", flush=True)
            elif event.type == "content_block_delta" and hasattr(event, "delta"):
                if event.delta.type == "text_delta":
                    print(".", end="", flush=True)
            elif event.type == "content_block_stop":
                if current == "text":
                    print("]", flush=True)
                current = None

        final_message = stream.get_final_message()

    text_content = "".join(
        block.text for block in final_message.content if block.type == "text"
    )

    try:
        json_match = re.search(r"\{[\s\S]*\}", text_content)
        return json.loads(json_match.group() if json_match else text_content)
    except json.JSONDecodeError as e:
        print(f"\nJSON パースエラー: {e}")
        print("レスポンス先頭500文字:", text_content[:500])
        raise


# ─── PowerPoint 生成 ──────────────────────────────────────────────────────────

def create_pptx(presentation_data: dict, output_path: str):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for slide_data in presentation_data.get("slides", []):
        if slide_data.get("type") == "title":
            _build_title_slide(prs, slide_data)
        else:
            _build_content_slide(prs, slide_data)

    prs.save(output_path)
    print(f"PowerPoint を保存しました: {output_path}")


# ─── エントリポイント ─────────────────────────────────────────────────────────

def main():
    transcription_file = sys.argv[1] if len(sys.argv) > 1 else "sample_transcription.txt"

    if not os.path.exists(transcription_file):
        print(f"エラー: ファイルが見つかりません: {transcription_file}")
        sys.exit(1)

    print(f"文字起こしファイル: {transcription_file}")
    with open(transcription_file, "r", encoding="utf-8") as f:
        transcription = f.read()

    presentation_data = research_and_structure(transcription)

    stem = Path(transcription_file).stem
    json_output = f"{stem}_structure.json"
    with open(json_output, "w", encoding="utf-8") as f:
        json.dump(presentation_data, f, ensure_ascii=False, indent=2)
    print(f"構成JSON を保存しました: {json_output}")

    pptx_output = f"{stem}_presentation.pptx"
    create_pptx(presentation_data, pptx_output)

    print(f"\n完了！生成ファイル:")
    print(f"  {json_output}  ← スライド構成（JSON）")
    print(f"  {pptx_output}  ← PowerPoint ファイル")


if __name__ == "__main__":
    main()
